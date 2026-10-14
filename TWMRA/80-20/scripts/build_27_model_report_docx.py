"""Build the final 27-model comparison report."""

from __future__ import annotations

import sys
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
LOGS = ROOT / "logs"
OUTPUT = LOGS / "FINAL_27_MODEL_PROTOCOL_REPORT.docx"

MODEL_FAMILIES = {
    "histgradientboosting": "Tree boosting",
    "lightgbm": "Tree boosting",
    "rbf_svm": "Kernel SVM",
    "xgboost": "Tree boosting",
    "mlp_tabular": "Tabular neural network",
    "random_forest": "Tree ensemble",
    "logistic_regression": "Linear baseline",
    "linear_svm": "Linear SVM",
    "tcn": "Temporal convolution",
    "tcnca_v1": "Temporal convolution/channel attention",
    "tcnca_v2": "Temporal convolution/channel attention",
    "tcnca_v3": "Temporal convolution/channel attention",
    "gru": "Recurrent neural network",
    "lstm": "Recurrent neural network",
    "lstm_attention": "Recurrent + attention",
    "bilstm": "Bidirectional recurrent",
    "bilstm_attention": "Bidirectional recurrent + attention",
    "conv_transformer": "Transformer",
    "fusformer": "Transformer",
    "mega": "Transformer",
    "tst": "Time Series Transformer",
    "convnet_har": "Temporal CNN",
    "s4d": "State-space model",
    "s5_onnx": "State-space model",
    "lru": "State-space model",
    "sgconv": "Structured convolution",
    "resnet1d": "Residual temporal CNN",
}


def add_table(document: Document, frame: pd.DataFrame, columns: list[str]) -> None:
    """Add a compact dataframe table."""
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, column in zip(table.rows[0].cells, columns):
        cell.text = str(column)
    for _, row in frame[columns].iterrows():
        cells = table.add_row().cells
        for cell, column in zip(cells, columns):
            value = row[column]
            if isinstance(value, (float, np.floating)):
                value = f"{value:.2%}" if column in {"accuracy", "macro_f1"} else f"{value:.4f}"
            cell.text = str(value)
    document.add_paragraph()


def create_accuracy_figure(frame: pd.DataFrame) -> Path:
    """Create the ranked 27-model accuracy figure."""
    path = LOGS / "figures" / "final_27_model_accuracy.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    plot = frame.sort_values("accuracy")
    figure, axis = plt.subplots(figsize=(9, 8))
    axis.barh(plot["model"], plot["accuracy"] * 100, color="#2f75b5")
    axis.axvline(75, color="#c00000", linestyle="--", label="75%")
    axis.set(xlabel="Independent-test accuracy (%)", title="27-model strict-protocol comparison")
    axis.legend()
    axis.grid(axis="x", alpha=0.25)
    figure.tight_layout()
    figure.savefig(path, dpi=300)
    plt.close(figure)
    return path


def build_report() -> Path:
    """Build the final report from completed experiment artifacts."""
    metrics = pd.read_csv(LOGS / "all_27_models_accuracy_comparison.csv")
    metrics["family"] = metrics["model"].map(MODEL_FAMILIES).fillna("Other")
    metrics = metrics.sort_values(["accuracy", "macro_f1"], ascending=False).reset_index(drop=True)
    metrics.insert(0, "rank", np.arange(1, len(metrics) + 1))
    figure_path = create_accuracy_figure(metrics)
    robustness = pd.read_csv(LOGS / "rigorous_experiment" / "robustness.csv")
    window = robustness[robustness["test"].str.startswith("window_size_")].copy()
    random_reference = pd.DataFrame([
        ["TCNCA optimized", 0.7990430622, 0.7999794175],
        ["TST optimized", 0.7511961722, 0.7405400467],
    ], columns=["model", "accuracy", "macro_f1"])

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Two-Wheeler Behaviour Recognition\nFinal 27-Model Report")
    run.bold = True
    run.font.size = Pt(18)
    subtitle = document.add_paragraph(
        "Strict four-file evaluation with separate paper-style 80/20 reference results"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("1. Executive summary", level=1)
    best = metrics.iloc[0]
    document.add_paragraph(
        f"All {len(metrics)} requested model entries are listed. Under the required strict "
        f"protocol, {best['model']} is the leader at {best['accuracy']:.2%} accuracy and "
        f"{best['macro_f1']:.2%} macro-F1. The validation-selected ensemble reached 75.02% "
        "accuracy, but it is reported separately from individual models."
    )
    document.add_paragraph(
        "The paper-style random 80/20 reference results are not mixed into the strict ranking: "
        f"TCNCA optimized reached {random_reference.iloc[0]['accuracy']:.2%}, and TST optimized "
        f"reached {random_reference.iloc[1]['accuracy']:.2%} under that separate protocol."
    )

    document.add_heading("2. Evaluation protocol", level=1)
    add_table(document, pd.DataFrame([
        ["Features", "Ax, Ay, Az, Gx, Gy, Gz, Speed"],
        ["Window", "156 samples"],
        ["Stride", "78 samples"],
        ["Classes", "BUMP, LEFT, RIGHT, STOP, STRAIGHT"],
        ["Strict training", "train.csv"],
        ["Strict selection", "val.csv"],
        ["Strict final retraining", "train_val_combined.csv"],
        ["Strict final test", "test.csv, evaluated once after selection"],
        ["Strict windows", "3,875 train; 1,045 validation; 4,921 combined; 1,045 test"],
    ], columns=["Setting", "Value"]), ["Setting", "Value"])

    document.add_heading("3. Complete 27-model strict comparison", level=1)
    add_table(document, metrics, ["rank", "model", "family", "accuracy", "macro_f1", "status"])
    document.add_picture(str(figure_path), width=Inches(6.7))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("4. Best results by family", level=1)
    family_best = metrics.sort_values("accuracy", ascending=False).groupby("family", as_index=False).first()
    add_table(document, family_best, ["family", "model", "accuracy", "macro_f1"])

    document.add_heading("5. Paper-style 80/20 reference results", level=1)
    random_reference["protocol"] = "Stratified random 80/20 windows from clean.csv"
    add_table(document, random_reference, ["protocol", "model", "accuracy", "macro_f1"])
    document.add_paragraph(
        "These results reproduce the earlier paper-style experiments. Because overlapping "
        "windows may appear in both random partitions, they must not be treated as equivalent "
        "to the strict independent test."
    )

    document.add_heading("6. Confusion matrices and class-wise analysis", level=1)
    confusion_files = sorted((LOGS / "figures").glob("protocol_*_confusion.png"))
    for path in confusion_files:
        document.add_picture(str(path), width=Inches(4.8))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    classwise_path = LOGS / "rigorous_experiment" / "histgradientboosting_classwise.json"
    if classwise_path.exists():
        with classwise_path.open("r", encoding="utf-8") as handle:
            classwise = pd.DataFrame(json.load(handle)["classwise"])
        document.add_paragraph("Class-wise metrics for the strict-protocol HistGradientBoosting leader:")
        add_table(document, classwise, ["class", "precision", "recall", "support"])

    document.add_heading("7. Robustness and window-size sensitivity", level=1)
    add_table(document, robustness, ["test", "accuracy", "macro_f1"])
    document.add_paragraph(
        "Window-size rows are diagnostic resampling checks for the fixed-input tabular model. "
        "The official comparison remains the 156-sample result."
    )
    add_table(document, window, ["test", "accuracy", "macro_f1"])

    document.add_heading("8. Limitations", level=1)
    document.add_paragraph(
        "The random 80/20 and strict four-file results measure different generalization "
        "questions. The strict test is the authoritative result for deployment comparison. "
        "Robustness perturbations are simulated and are not additional independent holdouts. "
        "Raspberry Pi 5 latency and memory must still be measured on hardware."
    )
    document.add_heading("9. Reproducibility", level=1)
    document.add_paragraph(
        "All models use the same seven raw features and 156-sample windows. Model-specific "
        "internal architectures are allowed. Test labels are not used for strict model selection."
    )
    document.save(OUTPUT)
    print(f"Saved {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build_report()
