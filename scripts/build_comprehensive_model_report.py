"""Build a comprehensive Word report for all implemented models."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import yaml
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sklearn.metrics import mean_absolute_error

ROOT = Path(__file__).resolve().parents[1]
LOGS = ROOT / "logs"
FIGURES = LOGS / "figures"
OUTPUT = ROOT / "COMPREHENSIVE_MODEL_REPORT_FINAL.docx"

MODEL_INFO: dict[str, dict[str, str]] = {
    "logistic_regression": {"family": "Linear baseline", "description": "Multiclass linear classifier over flattened 156x7 windows; interpretable and fast.", "hyperparameters": "max_iter=500, class_weight=balanced, solver=lbfgs"},
    "linear_svm": {"family": "Linear baseline", "description": "Linear maximum-margin classifier over flattened sensor windows.", "hyperparameters": "class_weight=balanced, max_iter=3000"},
    "rbf_svm": {"family": "Kernel baseline", "description": "Nonlinear RBF-kernel support-vector classifier for flattened windows.", "hyperparameters": "kernel=rbf, probability=True, class_weight=balanced"},
    "random_forest": {"family": "Tree ensemble", "description": "Bagged decision trees that model nonlinear feature interactions.", "hyperparameters": "n_estimators=100, class_weight=balanced, n_jobs=-1"},
    "histgradientboosting": {"family": "Tree boosting", "description": "Histogram-based gradient boosting over flattened sensor features.", "hyperparameters": "n_estimators-style boosting, max_iter=100, random_state=42"},
    "lightgbm": {"family": "Tree boosting", "description": "LightGBM multiclass gradient boosting optimized for tabular data.", "hyperparameters": "n_estimators=200, learning_rate=0.05, num_leaves=31, class_weight=balanced"},
    "xgboost": {"family": "Tree boosting", "description": "Regularized gradient-boosted trees used as the deployment baseline.", "hyperparameters": "n_estimators=200, max_depth=6, learning_rate=0.05, subsample=0.9, colsample_bytree=0.8"},
    "mlp_tabular": {"family": "Feed-forward neural network", "description": "Two-hidden-layer Keras MLP over flattened windows.", "hyperparameters": "Dense=128, Dense=64, dropout=0.2, Adam, batch_size=64, epochs=10"},
    "gru": {"family": "Recurrent neural network", "description": "Two-layer gated recurrent sequence model for temporal dependencies.", "hyperparameters": "GRU=64 return_sequences, dropout=0.2, GRU=32, Adam, batch_size=64, max_epochs=30"},
    "lstm": {"family": "Recurrent neural network", "description": "Single-layer LSTM HAR classifier.", "hyperparameters": "LSTM=64, dropout=0.2, Adam, batch_size=64, max_epochs=30"},
    "resnet1d": {"family": "Convolutional neural network", "description": "Residual 1D temporal convolution model exported to ONNX.", "hyperparameters": "input channels=7, base channels=16, residual blocks=2, Adam, epochs=10"},
    "convnet_har": {"family": "Convolutional neural network", "description": "Three-block 1D ConvNet with global average pooling.", "hyperparameters": "Conv1D filters=64, kernels=7/5/3, global average pooling, Adam, epochs=10"},
    "sgconv": {"family": "Structured convolution", "description": "Compact depthwise and dilated convolution approximation of SGConv.", "hyperparameters": "channels=32/64, dilation=2/4, GELU, Adam, epochs=10"},
    "s4d": {"family": "State-space model", "description": "Compact diagonal state-space recurrence exported to ONNX.", "hyperparameters": "state_dim=32, learned decay, LayerNorm, Adam, epochs=10"},
    "s5": {"family": "State-space model", "description": "Existing Keras/TFLite S5-style recurrent state-space classifier.", "hyperparameters": "state_size=16, sigmoid decay, Adam, epochs=12"},
    "s5_onnx": {"family": "State-space model", "description": "ONNX state-space variant evaluated alongside S4D and LRU.", "hyperparameters": "state_dim=32, learned decay, LayerNorm, Adam, epochs=10"},
    "lru": {"family": "State-space model", "description": "Linear recurrent unit approximation with learned phase modulation.", "hyperparameters": "state_dim=32, learned phase and decay, Adam, epochs=10"},
    "tst": {"family": "Transformer", "description": "Compact Temporal-Spatial Transformer with positional embeddings.", "hyperparameters": "d_model=24, heads=2, feedforward=48, blocks=2, dropout=0.1, epochs=12"},
    "tcnca_v1": {"family": "Temporal convolution ablation", "description": "Dilated temporal CNN without chunked attention.", "hyperparameters": "channels=24, dilations=[1,2,4,8,16], no attention, epochs=8"},
    "tcnca_v2": {"family": "Temporal convolution ablation", "description": "Dilated temporal CNN with chunked attention using 26-sample chunks.", "hyperparameters": "channels=24, chunk_size=26, dilations=[1,2,4,8,16], epochs=8"},
    "tcnca_v3": {"family": "Temporal convolution ablation", "description": "Final TCNCA ablation with larger chunks and five dilation rates.", "hyperparameters": "channels=24, chunk_size=52, dilations=[1,2,4,8,16], epochs=8"},
    "mega": {"family": "Transformer/attention", "description": "Moving-average gated attention approximation for temporal smoothing and attention.", "hyperparameters": "Conv1D channels=32, average pool=9, heads=4, key_dim=8, epochs=8"},
    "fusformer": {"family": "Transformer/attention", "description": "Fusion Transformer combining short and long convolutional branches.", "hyperparameters": "Conv kernels=5/11, fused channels=32, heads=4, epochs=8"},
    "conv_transformer": {"family": "Transformer/attention", "description": "Convolutional stem followed by two transformer encoder blocks.", "hyperparameters": "Conv1D channels=32, heads=4, blocks=2, feedforward=64, epochs=8"},
    "dtaad": {"family": "Anomaly detection", "description": "Dual dilated-TCN autoencoder with temporal attention; returns novelty score, not a class.", "hyperparameters": "Two TCN branches, dilations=1/2/4, attention heads=2, normal label=STRAIGHT, epochs=15"},
    "egodrive_rt": {"family": "Stub/future", "description": "Stub only because pretrained EgoDriveRT weights are not public.", "hyperparameters": "Not applicable"},
    "egodrive_max": {"family": "Stub/rejected", "description": "Rejected because the model is expected to exceed 10M parameters and the 10 ms budget.", "hyperparameters": "Not applicable"},
}


def set_cell_shading(cell: object, fill: str) -> None:
    """Apply background shading to a Word table cell."""
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table(document: Document, frame: pd.DataFrame, columns: list[str], headers: list[str] | None = None) -> None:
    """Add a dataframe as a formatted Word table."""
    if frame.empty:
        document.add_paragraph("No data available.")
        return
    headers = headers or columns
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_shading(cell, "D9EAF7")
    for _, row in frame[columns].iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            value = row[column]
            if isinstance(value, (float, np.floating)):
                value = f"{value:.4f}"
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8)
    document.add_paragraph()


def add_picture(document: Document, path: Path, caption: str, width: float = 6.3) -> None:
    """Embed an existing image with a caption."""
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.runs[0].italic = True
    caption_paragraph.runs[0].font.size = Pt(8)


def load_config() -> dict[str, Any]:
    """Load project configuration."""
    return yaml.safe_load((ROOT / "config.yaml").read_text(encoding="utf-8"))


def predict_for_model(model_name: str, X: np.ndarray, labels: list[str], config: dict[str, Any]) -> np.ndarray:
    """Generate integer predictions using the saved adapter artifact."""
    from scripts.benchmark_all import artifact_path, build_adapter

    artifact = artifact_path(model_name, ROOT / config["model"]["models_dir"])
    adapter = build_adapter(model_name, artifact, len(labels), labels)
    return np.asarray([int(np.argmax(adapter.predict(window))) for window in X], dtype=np.int64)


def calculate_mae(benchmark: pd.DataFrame, config: dict[str, Any]) -> pd.DataFrame:
    """Calculate encoded-label MAE for each classification model."""
    with np.load(ROOT / config["data"]["windows_npz"], allow_pickle=True) as data:
        X = np.asarray(data["X"], dtype=np.float32)[:500]
        y = np.asarray(data["y"]).astype(str)[:500]
    labels = sorted(np.unique(y).tolist())
    mapping = {label: index for index, label in enumerate(labels)}
    y_true = np.asarray([mapping[label] for label in y], dtype=np.int64)
    rows: list[dict[str, object]] = []
    for model_name in benchmark.loc[benchmark["status"] == "benchmarked", "model"]:
        if model_name not in MODEL_INFO or model_name == "dtaad":
            continue
        try:
            predictions = predict_for_model(model_name, X, labels, config)
            rows.append({"model": model_name, "mae": float(mean_absolute_error(y_true, predictions))})
        except Exception as exc:
            rows.append({"model": model_name, "mae": np.nan, "mae_error": str(exc)})
    result = pd.DataFrame(rows)
    result.to_csv(LOGS / "model_mae.csv", index=False)
    return result


def build_report() -> Path:
    """Build the comprehensive model comparison Word report."""
    config = load_config()
    benchmark = pd.read_csv(LOGS / "benchmark_all.csv")
    mae_path = LOGS / "model_mae.csv"
    mae = pd.read_csv(mae_path) if mae_path.exists() else calculate_mae(benchmark, config)
    mae_by_model = dict(zip(mae["model"], mae["mae"]))
    benchmark["mae"] = benchmark["model"].map(mae_by_model)
    benchmark["description"] = benchmark["model"].map(lambda name: MODEL_INFO.get(name, {}).get("description", "Not documented"))
    benchmark["family"] = benchmark["model"].map(lambda name: MODEL_INFO.get(name, {}).get("family", "Other"))

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    document.styles["Title"].font.color.rgb = RGBColor(31, 78, 121)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Two-Wheeler Driving Behaviour Recognition\nComprehensive Model Report")
    subtitle = document.add_paragraph("Model explanations, hyperparameters, metrics, graphs, confusion matrices, and final comparison")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()

    document.add_heading("1. Executive Summary", level=1)
    best = benchmark[benchmark["status"] == "benchmarked"].sort_values("macro_f1", ascending=False).iloc[0]
    document.add_paragraph(f"The strongest measured classification result is {best['model']} with accuracy {best['accuracy']:.2%} and macro F1 {best['macro_f1']:.4f}. The current benchmark evaluates 500 windows from the full window artifact because only 209 random-test windows are available; therefore these values are comparative engineering measurements, not independent generalisation estimates.")
    document.add_paragraph("MAE is reported on integer-encoded labels as a diagnostic only. Since the behaviour classes are nominal, macro F1 is the primary balanced metric.")

    document.add_heading("2. Dataset and Metrics", level=1)
    dataset = pd.DataFrame([
        ["Features", "Ax, Ay, Az, Gx, Gy, Gz, Speed"],
        ["Window", "156 samples"],
        ["Stride", "78 samples"],
        ["Classes", "STRAIGHT, LEFT, RIGHT, BUMP, STOP"],
        ["Accuracy", "Correct predictions / total predictions"],
        ["Macro F1", "Unweighted mean of the five class F1 scores"],
        ["MAE", "Mean absolute error on encoded class IDs; diagnostic only"],
    ], columns=["Item", "Value"])
    add_table(document, dataset, ["Item", "Value"])

    document.add_heading("3. Model Explanations and Hyperparameters", level=1)
    explanations = []
    for name, info in MODEL_INFO.items():
        explanations.append({"model": name, "family": info["family"], "explanation": info["description"], "hyperparameters": info["hyperparameters"]})
    add_table(document, pd.DataFrame(explanations), ["model", "family", "explanation", "hyperparameters"], ["Model", "Family", "What it does", "Hyperparameters"])

    document.add_heading("4. Metrics for Every Model", level=1)
    metrics_columns = ["model", "family", "mae", "accuracy", "macro_f1", "latency_mean_ms", "latency_p95_ms", "params", "status"]
    metrics_table = benchmark[metrics_columns].sort_values("macro_f1", ascending=False, na_position="last").copy()
    metrics_table["mae"] = metrics_table["mae"].map(lambda value: "N/A" if pd.isna(value) else f"{float(value):.4f}")
    add_table(document, metrics_table, metrics_columns, ["Model", "Family", "MAE", "Accuracy", "Macro F1", "Mean latency (ms)", "P95 latency (ms)", "Parameters", "Status"])
    document.add_paragraph("DTAAD is anomaly-only and therefore has no classification accuracy, F1, or class-label MAE.")

    document.add_heading("5. Benchmark Graphs", level=1)
    for name, caption in [("accuracy.png", "Accuracy comparison"), ("macro_f1.png", "Macro F1 comparison"), ("mae.png", "Encoded-label MAE comparison"), ("latency_ms.png", "Latency comparison"), ("optimized_neural_comparison.png", "Optimized neural comparison")]:
        add_picture(document, FIGURES / name, caption)

    document.add_heading("6. Confusion Matrices and Class-Wise Results", level=1)
    classwise_path = LOGS / "classwise_all.csv"
    if classwise_path.exists():
        classwise = pd.read_csv(classwise_path)
        add_table(document, classwise, ["model", "class", "precision", "recall", "f1", "support"], ["Model", "Class", "Precision", "Recall", "F1", "Support"])
    for path in sorted(FIGURES.glob("cm_*.png")):
        add_picture(document, path, f"Confusion matrix: {path.stem.removeprefix('cm_')}", width=5.6)

    document.add_heading("7. TCNCA Ablation and Robustness", level=1)
    tcnca_path = LOGS / "benchmark_transformers.csv"
    if tcnca_path.exists():
        tcnca = pd.read_csv(tcnca_path)
        add_table(document, tcnca[tcnca["model"].str.startswith("tcnca")], ["model", "accuracy", "macro_f1", "latency_ms", "params"], ["Variant", "Accuracy", "Macro F1", "Latency (ms)", "Parameters"])
    for path in sorted(LOGS.glob("robustness_*.csv")):
        document.add_heading(path.stem.replace("_", " ").title(), level=2)
        add_table(document, pd.read_csv(path), list(pd.read_csv(path).columns))
        add_picture(document, FIGURES / f"{path.stem}.png", f"{path.stem.replace('_', ' ').title()} summary")

    document.add_heading("8. Final Comparison and Recommendation", level=1)
    comparison = benchmark[benchmark["status"] == "benchmarked"].sort_values("macro_f1", ascending=False)[["model", "accuracy", "macro_f1", "mae", "latency_p95_ms"]]
    comparison = comparison.copy()
    comparison["mae"] = comparison["mae"].map(lambda value: "N/A" if pd.isna(value) else f"{float(value):.4f}")
    add_table(document, comparison, list(comparison.columns), ["Model", "Accuracy", "Macro F1", "MAE", "P95 latency (ms)"])
    document.add_paragraph(f"Recommendation: use XGBoost as the production classification baseline because it combines {float(benchmark.loc[benchmark['model'].eq('xgboost'), 'accuracy'].iloc[0]):.2%} accuracy, strong macro F1, and low measured latency. HistGradientBoosting has the highest score in the 500-window comparison, but its p95 latency is materially higher. ConvTransformer is the strongest compact neural candidate among the measured transformer models. Final deployment selection requires Raspberry Pi 5 measurements.")

    document.add_heading("9. Limitations", level=1)
    for limitation in [
        "Only two sessions are available; generalisation to unseen rides or riders cannot be claimed.",
        "The unified 500-window benchmark uses the full window artifact because the random test split has only 209 windows, so training-data reuse is possible.",
        "Session-wise comparisons reuse existing random-split-trained artifacts and are diagnostic rather than clean retrain-per-session validation.",
        "MAE on encoded labels is not semantically ordered and should not replace macro F1.",
        "Clock-jitter experiments are simulated, not real field clock drift measurements.",
        "EgoDriveRT weights are unavailable and EgoDriveMax is rejected for deployment constraints.",
    ]:
        document.add_paragraph(limitation, style="List Bullet")

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Two-Wheeler Driving Behaviour Recognition | Comprehensive Model Report").font.size = Pt(8)
    document.save(OUTPUT)
    print(f"Saved Word report: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build_report()
