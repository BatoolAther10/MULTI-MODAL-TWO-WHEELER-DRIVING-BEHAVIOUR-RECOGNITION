"""Build the final two-wheeler model report as a Word document."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
LOGS = ROOT / "logs"
FIGURES = LOGS / "figures"
TUNING_FIGURES = FIGURES / "tuning"
OUTPUT = ROOT / "FINAL_REPORT.docx"

MODEL_DESCRIPTIONS = {
    "XGBoost": (
        "Gradient-boosted decision-tree classifier on flattened sensor windows.",
        "It is a strong non-neural baseline for tabular sensor features, handles nonlinear interactions, trains offline, and is fast enough for the 10 ms fallback/primary requirement.",
        "n_estimators, max_depth, learning_rate, subsample, colsample_bytree"
    ),
    "GRU": (
        "Gated recurrent neural network with two recurrent layers.",
        "It models temporal dependencies and is a natural sequence-learning baseline for inertial time series, but recurrent TensorFlow execution is slow in this environment.",
        "GRU widths 32/16 or 64/32, learning rate, epochs, batch size"
    ),
    "ResNet1D": (
        "One-dimensional residual convolutional neural network exported to ONNX.",
        "Residual temporal convolutions can learn local motion patterns with low inference cost and provide a CNN comparison to recurrent and attention models.",
        "channels, residual block count, learning rate, epochs"
    ),
    "TST": (
        "Compact Temporal-Spatial Transformer using self-attention, feed-forward residual blocks, normalization, and positional information.",
        "Self-attention can capture relationships across the entire riding window and is a candidate primary model for multimodal temporal data.",
        "d_model, attention heads, feed-forward width, block count, dropout, learning rate"
    ),
    "TCNCA": (
        "Temporal convolutional network with dilated convolutions, residual connections, and channel attention.",
        "It combines fast convolutional inference with attention over sensor channels, making it a practical neural model for Raspberry Pi deployment.",
        "channels, dilation/residual block count, learning rate, epochs"
    ),
    "S5": (
        "Compact state-space sequence model with a learned recurrent state update.",
        "State-space recurrence is included as a long-context alternative to GRU and Transformer models with a small parameter count.",
        "state size, decay parameters, learning rate, epochs"
    ),
}


def set_cell_shading(cell, fill: str) -> None:
    """Apply a background color to a table cell."""
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    """Replace cell contents with formatted text."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, dataframe: pd.DataFrame, columns: list[str], headers: list[str] | None = None) -> None:
    """Add a compact dataframe table to the document."""
    headers = headers or columns
    table = document.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], "D9EAF7")
    for _, row in dataframe[columns].iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            value = row[column]
            if isinstance(value, float):
                value = f"{value:.3f}"
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_figure(document: Document, path: Path, caption: str, width: float = 6.3) -> None:
    """Embed a figure with a caption when it exists."""
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.runs[0].italic = True
    caption_paragraph.runs[0].font.size = Pt(9)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    """Add a heading with consistent report styling."""
    document.add_heading(text, level=level)


def format_metric(value: object) -> object:
    """Format numeric report values compactly."""
    if isinstance(value, float):
        return round(value, 4)
    return value


def build_report() -> Path:
    """Build and save the final report."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)

    title = document.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Two-Wheeler Driving Behaviour Recognition\nFinal Model and Accuracy Report")
    subtitle = document.add_paragraph("Offline multimodal sensor classification pipeline for Raspberry Pi 5")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    document.add_paragraph()

    add_heading(document, "1. Executive Summary")
    document.add_paragraph(
        "This report documents the complete two-wheeler driving-behaviour classification pipeline. "
        "The system converts IMU/GPS sensor records into fixed temporal windows, trains several classification model families, tunes their principal hyperparameters, calculates accuracy, macro F1, and encoded-label mean absolute error (MAE), and compares deployment latency."
    )
    document.add_paragraph(
        "The verified benchmark leader is XGBoost with 94.0% accuracy, 94.1% macro F1, and approximately 3.40 ms median inference latency. "
        "The strongest neural candidate is the optimized TCNCA model at approximately 79.9% accuracy and 80.0% macro F1 on the same holdout."
    )

    add_heading(document, "2. Dataset and Learning Task")
    document.add_paragraph("Input file: data/raw/test.csv. The cleaned and windowed artifacts are data/processed/clean.csv and data/processed/windows.npz.")
    dataset_table = pd.DataFrame([
        ["Input features", "Ax, Ay, Az, Gx, Gy, Gz, Speed"],
        ["Window", "156 samples x 7 features"],
        ["Stride", "78 samples"],
        ["Classes", "STRAIGHT, LEFT, RIGHT, BUMP, STOP"],
        ["Learning task", "Supervised multiclass classification"],
        ["Validation split", "80/20 stratified random holdout"],
    ], columns=["Item", "Value"])
    add_table(document, dataset_table, ["Item", "Value"])
    document.add_paragraph(
        "Each window receives one majority-vote behaviour label. The current evaluation is a local holdout comparison; a rider/session-based split is recommended before claiming production generalization."
    )

    add_heading(document, "3. What Was Learned and Why Each Model Was Used")
    for name, (model_type, reason, hyperparameters) in MODEL_DESCRIPTIONS.items():
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{name}: ").bold = True
        paragraph.add_run(f"Model type: {model_type} Learning rationale: {reason} Main hyperparameters: {hyperparameters}.")

    add_heading(document, "4. Classification Metrics and MAE Calculation")
    document.add_paragraph(
        "Accuracy is the fraction of correctly classified windows. Macro F1 computes F1 independently for each class and averages the class scores, giving minority classes equal importance."
    )
    document.add_paragraph(
        "MAE was calculated on integer-encoded class labels using: MAE = mean(abs(y_true_encoded - y_pred_encoded)). "
        "This is a diagnostic metric only because class IDs are nominal rather than ordered; macro F1 is the preferred class-balanced metric."
    )
    document.add_paragraph(
        "For example, with encoded labels BUMP=0, LEFT=1, RIGHT=2, STOP=3, STRAIGHT=4, a prediction of 3 instead of 1 contributes an absolute error of 2."
    )

    benchmark = pd.read_csv(LOGS / "benchmark.csv")
    benchmark["accuracy"] = benchmark["accuracy"].map(format_metric)
    benchmark["macro_f1"] = benchmark["macro_f1"].map(format_metric)
    benchmark["latency_ms"] = benchmark["latency_ms"].map(format_metric)
    add_heading(document, "5. All-Model Benchmark Comparison")
    add_table(document, benchmark, ["model", "accuracy", "macro_f1", "latency_ms"], ["Model", "Accuracy", "Macro F1", "Latency (ms)"])
    add_figure(document, FIGURES / "accuracy.png", "Figure 1. Accuracy comparison using the benchmark artifact.")
    add_figure(document, FIGURES / "macro_f1.png", "Figure 2. Macro F1 comparison using the benchmark artifact.")
    add_figure(document, FIGURES / "latency_ms.png", "Figure 3. Measured deployment latency comparison.")
    add_figure(document, FIGURES / "mae.png", "Figure 4. Encoded-label MAE comparison; baseline rows without saved MAE are omitted.")
    add_figure(document, FIGURES / "optimized_neural_comparison.png", "Figure 5. Latest optimized TCNCA and TST comparison.")

    add_heading(document, "6. Hyperparameters and Tuning Results")
    document.add_paragraph(
        "The tuning process varied architecture size and optimization settings using a fixed stratified holdout. Each model family received a compact sweep appropriate to its computational cost. The selected configuration was the strongest trial by macro F1, with accuracy used as a secondary criterion."
    )
    tuning_files = {
        "xgboost": "xgboost_tuning.csv",
        "tcnca": "tcnca_tuning.csv",
        "tst": "tst_tuning.csv",
        "resnet1d": "resnet1d_tuning.csv",
        "gru": "gru_tuning.csv",
        "s5": "s5_tuning.csv",
    }
    best_rows = []
    for model, filename in tuning_files.items():
        path = LOGS / filename
        if not path.exists():
            continue
        data = pd.read_csv(path).sort_values(["macro_f1", "accuracy"], ascending=False)
        best = data.iloc[0]
        best_rows.append({
            "model": model,
            "accuracy": best["accuracy"],
            "macro_f1": best["macro_f1"],
            "mae": best["mae"],
            "latency_ms": best["latency_ms"],
            "num_params": best.get("num_params", ""),
        })
    best_table = pd.DataFrame(best_rows)
    add_table(document, best_table, ["model", "accuracy", "macro_f1", "mae", "latency_ms", "num_params"], ["Model", "Accuracy", "Macro F1", "MAE", "Tuning latency (ms)", "Parameters"])
    add_figure(document, TUNING_FIGURES / "tuning_summary.png", "Figure 6. Best tuned configuration comparison.")
    for model in tuning_files:
        add_figure(document, TUNING_FIGURES / f"{model}_tuning.png", f"Figure. {model.upper()} hyperparameter variation: accuracy and MAE.")

    add_heading(document, "7. Optimized Neural Training")
    optimized_rows = []
    for model in ("tcnca", "tst"):
        path = LOGS / f"{model}_optimized_metrics.json"
        if path.exists():
            metrics = json.loads(path.read_text(encoding="utf-8"))
            optimized_rows.append({
                "model": model,
                "accuracy": metrics["accuracy"],
                "macro_f1": metrics["macro_f1"],
                "mae": metrics["mae"],
                "parameters": metrics["num_params"],
                "configuration": json.dumps(metrics["configuration"]),
            })
    optimized_table = pd.DataFrame(optimized_rows)
    add_table(document, optimized_table, ["model", "accuracy", "macro_f1", "mae", "parameters", "configuration"], ["Model", "Accuracy", "Macro F1", "MAE", "Parameters", "Configuration"])
    document.add_paragraph(
        "The optimized neural run added feature normalization inside the exported model, positional encoding for TST, longer early-stopped training, and a wider TCNCA configuration. "
        "TCNCA reached 79.9% accuracy and TST reached 75.1% accuracy on the unchanged holdout."
    )

    add_heading(document, "8. Deployment and Final Recommendation")
    document.add_paragraph(
        "The runtime uses XGBoost as the primary model because it has the highest verified accuracy and remains within the 10 ms budget. "
        "Optimized TCNCA is configured as the neural fallback. If the primary model raises an exception or exceeds the latency budget, the runtime invokes the fallback model."
    )
    deployment = pd.DataFrame([
        ["Primary", "XGBoost", "Highest verified benchmark accuracy/F1; approximately 3.40 ms median benchmark latency"],
        ["Fallback", "Optimized TCNCA", "Best neural candidate; approximately 79.9% holdout accuracy and fast TFLite path"],
        ["Budget", "10 ms", "Configured Raspberry Pi inference constraint"],
    ], columns=["Role", "Model", "Reason"])
    add_table(document, deployment, ["Role", "Model", "Reason"])
    document.add_paragraph(
        "The runtime smoke test successfully loaded the configured primary/fallback path and produced a real-window prediction. Reported latency can vary by machine load; Raspberry Pi validation remains necessary."
    )

    add_heading(document, "9. Limitations and Future Work")
    limitations = [
        "The random stratified holdout uses overlapping windows, so nearby windows may be correlated. Use session-based or rider-based splits for final claims.",
        "Encoded-label MAE is diagnostic because behaviour classes are nominal.",
        "GRU and S5 require fallback handling in this TensorFlow environment because their exported graphs contain Select TF operations.",
        "The 80-90% target is achieved by XGBoost on the current benchmark, while optimized TCNCA is close to 80% and optimized TST is approximately 75%.",
        "Additional labeled rides, better session separation, feature engineering, and real Raspberry Pi measurements are the next steps for stronger generalization.",
    ]
    for item in limitations:
        document.add_paragraph(item, style="List Bullet")

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Two-Wheeler Driving Behaviour Recognition | Final Report").font.size = Pt(8)

    document.save(OUTPUT)
    print(f"Saved Word report: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build_report()
